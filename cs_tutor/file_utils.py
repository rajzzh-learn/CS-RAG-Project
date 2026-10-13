"""
Utility helpers for processing file attachments uploaded via the Streamlit UI.

Supported attachment types:
  - PDF               → extracted plain text via pypdf
  - DOCX              → paragraphs + tables via python-docx
  - PPTX              → slide text + tables via python-pptx
  - XLSX              → sheet cells via openpyxl (read-only)
  - Text-like files   → raw decoded text (txt, md, csv, json, py, sql, html, ...)
  - Image (PNG / JPG / JPEG / WEBP / GIF) → base64-encoded data URI for vision LLMs
"""
from __future__ import annotations

import base64
import io
import re
from typing import Tuple

# Hard cap on a single attachment (also enforced by server.maxUploadSize)
MAX_FILE_BYTES = 10 * 1024 * 1024


def _truncate(text: str, max_chars: int) -> str:
    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n… [truncated at {max_chars} chars]"
    return text


# ── PDF extraction ──────────────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes, max_chars: int = 6000) -> str:
    """Extract plain text from PDF bytes using pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text: list[str] = []
    for page in reader.pages:
        pages_text.append((page.extract_text() or "").strip())
        if sum(len(t) for t in pages_text) > max_chars:
            break
    return _truncate("\n\n".join(filter(None, pages_text)), max_chars)


# ── Plain-text extraction ───────────────────────────────────────────────────

def extract_text_from_txt(file_bytes: bytes, max_chars: int = 6000) -> str:
    """Decode a plain-text file, truncating to *max_chars*."""
    return _truncate(file_bytes.decode("utf-8", errors="replace"), max_chars)


# ── DOCX extraction ─────────────────────────────────────────────────────────

def extract_text_from_docx(file_bytes: bytes, max_chars: int = 6000) -> str:
    """Extract paragraphs and table rows from a .docx file."""
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _truncate("\n".join(parts), max_chars)


# ── PPTX extraction ─────────────────────────────────────────────────────────

def extract_text_from_pptx(file_bytes: bytes, max_chars: int = 6000) -> str:
    """Extract text frames and table cells from a .pptx file."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {idx} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
    return _truncate("\n".join(parts), max_chars)


# ── XLSX extraction ─────────────────────────────────────────────────────────

def extract_text_from_xlsx(file_bytes: bytes, max_chars: int = 6000) -> str:
    """Extract cell values from every sheet of an .xlsx file."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []
    total = 0
    try:
        for ws in wb.worksheets:
            parts.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                line = " | ".join("" if v is None else str(v) for v in row).strip(" |")
                if line:
                    parts.append(line)
                    total += len(line)
                if total > max_chars:
                    break
            if total > max_chars:
                break
    finally:
        wb.close()
    return _truncate("\n".join(parts), max_chars)


# ── File type registry ──────────────────────────────────────────────────────

_MIME_MAP = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
    "gif": "image/gif",
}

SUPPORTED_IMAGE_EXTS = set(_MIME_MAP.keys())
TEXT_LIKE_EXTS = {"txt", "md", "csv", "json", "log", "py", "sql", "html", "htm", "xml"}
_EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "pptx": extract_text_from_pptx,
    "xlsx": extract_text_from_xlsx,
}
SUPPORTED_DOC_EXTS = set(_EXTRACTORS) | TEXT_LIKE_EXTS
SUPPORTED_EXTS = SUPPORTED_IMAGE_EXTS | SUPPORTED_DOC_EXTS


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def is_image(filename: str) -> bool:
    return _ext(filename) in SUPPORTED_IMAGE_EXTS


def extract_text(file_bytes: bytes, filename: str, max_chars: int = 6000) -> str:
    """Extract text from any supported non-image attachment."""
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ValueError(f"'{filename}' exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB limit.")
    ext = _ext(filename)
    extractor = _EXTRACTORS.get(ext)
    if extractor:
        return extractor(file_bytes, max_chars)
    if ext in TEXT_LIKE_EXTS:
        return extract_text_from_txt(file_bytes, max_chars)
    raise ValueError(f"Unsupported file type: '{filename}'")


def image_to_base64_uri(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Convert raw image bytes to a base64 data URI.

    Returns:
        (data_uri, mime_type)  e.g. ("data:image/png;base64,…", "image/png")
    """
    mime = _MIME_MAP.get(_ext(filename), "image/png")
    b64 = base64.b64encode(file_bytes).decode("ascii")
    return f"data:{mime};base64,{b64}", mime


# ── Rendering helpers (math + HTML) ─────────────────────────────────────────

_CODE_SPLIT = re.compile(r"(```.*?```|`[^`\n]*`)", re.DOTALL)


def normalize_math(text: str) -> str:
    """
    Convert LaTeX delimiters \\( \\) and \\[ \\] (which Streamlit/KaTeX does not
    parse) into $ … $ and $$ … $$. Code blocks and inline code are left untouched.
    """
    def _convert(segment: str) -> str:
        segment = re.sub(r"\\\[(.+?)\\\]", lambda m: f"\n$$\n{m.group(1).strip()}\n$$\n", segment, flags=re.DOTALL)
        segment = re.sub(r"\\\((.+?)\\\)", lambda m: f"${m.group(1).strip()}$", segment, flags=re.DOTALL)
        return segment

    parts = _CODE_SPLIT.split(text)
    # re.split with a capture group puts code segments at odd indices
    return "".join(p if i % 2 else _convert(p) for i, p in enumerate(parts))


_HTML_BLOCK = re.compile(r"```html[ \t]*\n(.*?)```", re.DOTALL | re.IGNORECASE)


def split_html_blocks(text: str) -> list[tuple[str, str]]:
    """
    Split markdown into [("md", text), ("html", source), …] so ```html blocks
    can be rendered as live previews.
    """
    out: list[tuple[str, str]] = []
    pos = 0
    for m in _HTML_BLOCK.finditer(text):
        if m.start() > pos:
            out.append(("md", text[pos:m.start()]))
        out.append(("html", m.group(1)))
        pos = m.end()
    if pos < len(text):
        out.append(("md", text[pos:]))
    return out


def sanitize_html(html: str) -> str:
    """
    Sanitize untrusted HTML with nh3 (allowlist-based): scripts, iframes, forms,
    event handlers and javascript:/data: URLs are removed.
    """
    import nh3

    return nh3.clean(html)
